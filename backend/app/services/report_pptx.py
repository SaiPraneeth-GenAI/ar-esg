"""Renders ReportSection/ReportMeta (services/dashboard_report.py) to PPTX
bytes via python-pptx -- plain layout code, no AI, no network calls."""

import io
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.services.dashboard_report import ReportMeta, ReportSection

_INK = RGBColor(0x10, 0x18, 0x28)
_GREEN = RGBColor(0x0E, 0x2C, 0x21)
_MUTED = RGBColor(0x66, 0x70, 0x85)

_LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "amara_raja_logo.png"
_LOGO_ASPECT = 202 / 682  # height / width of the source lockup


def build_pptx(sections: list[ReportSection], meta: ReportMeta) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    _add_title_slide(prs, blank_layout, meta)
    for i, section in enumerate(sections):
        _add_section_slide(prs, blank_layout, section, i + 2)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_footer(slide, prs: Presentation, slide_number: int) -> None:
    """Amara Raja branding on every slide -- what makes an exported deck
    look like it came from a real tenant deployment, not a generic tool."""
    if _LOGO_PATH.exists():
        logo_w = Inches(1.05)
        slide.shapes.add_picture(
            str(_LOGO_PATH), Inches(0.4), prs.slide_height - Inches(0.5), width=logo_w, height=logo_w * _LOGO_ASPECT
        )

    footer_box = slide.shapes.add_textbox(prs.slide_width - Inches(3.6), prs.slide_height - Inches(0.45), Inches(3.2), Inches(0.35))
    p = footer_box.text_frame.paragraphs[0]
    p.text = f"Slide {slide_number} · Generated via Enviqo"
    p.font.size = Pt(9)
    p.font.color.rgb = _MUTED
    p.alignment = PP_ALIGN.RIGHT


def _add_title_slide(prs: Presentation, layout, meta: ReportMeta) -> None:
    slide = prs.slides.add_slide(layout)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(11.7), Inches(1.3))
    p = title_box.text_frame.paragraphs[0]
    p.text = meta.title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = _INK

    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.6))
    p = sub_box.text_frame.paragraphs[0]
    p.text = meta.subtitle
    p.font.size = Pt(18)
    p.font.color.rgb = _MUTED

    _add_footer(slide, prs, 1)


def _add_section_slide(prs: Presentation, layout, section: ReportSection, slide_number: int) -> None:
    slide = prs.slides.add_slide(layout)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7))
    p = title_box.text_frame.paragraphs[0]
    p.text = section.title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = _GREEN

    narrative_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(12.3), Inches(1.0))
    narrative_box.text_frame.word_wrap = True
    p = narrative_box.text_frame.paragraphs[0]
    p.text = section.narrative
    p.font.size = Pt(13)
    p.font.color.rgb = _INK

    tiles_top = Inches(2.15)
    if section.tiles:
        tile_box = slide.shapes.add_textbox(Inches(0.5), tiles_top, Inches(5.2), Inches(4.8))
        tf = tile_box.text_frame
        tf.word_wrap = True
        for i, tile in enumerate(section.tiles):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"{tile.label}:  {tile.value}"
            p.font.size = Pt(15)
            p.font.color.rgb = _INK
            p.space_after = Pt(10)

    if section.chart_png:
        slide.shapes.add_picture(io.BytesIO(section.chart_png), Inches(6.0), tiles_top, width=Inches(6.8))

    _add_footer(slide, prs, slide_number)
